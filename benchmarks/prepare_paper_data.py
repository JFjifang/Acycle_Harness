"""Recreate lossless paper-fixture extraction; no scientific computation.

Requires openpyxl only for workbook parsing. Publisher originals are downloaded
only with --download; existing files are validated against downloads.json.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
from pathlib import Path
import urllib.request
from zipfile import ZipFile


ROOT = Path(__file__).resolve().parent


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def extract(source_name, output_name, sheet, columns, labels, doi, scope,
            member=None, bounds=None):
    from openpyxl import load_workbook
    source = ROOT / 'sources' / source_name
    data = source.read_bytes()
    archive_hash = sha256(data)
    if member:
        with ZipFile(io.BytesIO(data)) as archive:
            data = archive.read(member)
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows, source_rows = [], []
    for index, row in enumerate(workbook[sheet].iter_rows(values_only=True), 1):
        x, y = (row[column] for column in columns)
        if not isinstance(x, (int, float)) or not isinstance(y, (int, float)):
            continue
        if bounds and not bounds[0] <= x <= bounds[1]:
            continue
        rows.append((x, y))
        source_rows.append(index)
    workbook.close()
    output = ROOT / 'data' / output_name
    output.parent.mkdir(exist_ok=True)
    with output.open('w', newline='', encoding='utf-8') as stream:
        writer = csv.writer(stream)
        writer.writerow(labels)
        writer.writerows(rows)
    receipt = {
        'paper_doi': doi,
        'source_file': str(source.relative_to(ROOT)),
        'source_sha256': archive_hash,
        'sheet': sheet,
        'coordinate_column': chr(65 + columns[0]),
        'value_column': chr(65 + columns[1]),
        'selection': f'{bounds[0]} <= depth_m <= {bounds[1]}; numeric cells only; source row order unchanged'
        if bounds else 'All numeric rows, source order unchanged',
        'source_row_numbers': source_rows,
        'output_file': str(output.relative_to(ROOT)),
        'output_sha256': sha256(output.read_bytes()),
        'row_count': len(rows),
        'transformations': 'None. Python only extracts workbook cells and selects an explicit interval where specified; no resampling, detrending, spectral calculation or plot.',
        'scope': scope,
    }
    if member:
        receipt.update(archive_member=member, member_sha256=sha256(data))
    output.with_suffix('.provenance.json').write_text(
        json.dumps(receipt, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    return receipt


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--download', action='store_true')
    args = parser.parse_args()
    for record in json.loads((ROOT / 'sources/downloads.json').read_text()):
        path = ROOT / record['file']
        if not path.is_file():
            if not args.download:
                raise SystemExit(f'Missing {path}; use --download to fetch the recorded publisher source')
            with urllib.request.urlopen(record['source_url'], timeout=120) as response:
                data = response.read()
            if sha256(data) != record['sha256']:
                raise SystemExit(f'Publisher content changed: {path.name}')
            path.write_bytes(data)
        if sha256(path.read_bytes()) != record['sha256']:
            raise SystemExit(f'Source checksum mismatch: {path}')
    records = [
        extract('petm_2022_source.xlsx', 'petm_2022_ms_194_200m.csv', 'Ca & MS',
                (3, 4), ['depth_m', 'magnetic_susceptibility_SI_x1e-5'],
                '10.1038/s41467-022-33390-x',
                'Paper-data workflow smoke test on a selected 6 m interval, not reproduction of the published Lomb-Scargle/COCO/ASM analysis.',
                bounds=(194, 200)),
        extract('cambrian_2022_source.zip', 'cambrian_2022_gamma_ray.csv',
                'GR_Albjära-1 core', (0, 1), ['depth_m', 'gamma_ray_counts_per_second'],
                '10.1038/s41467-022-29651-4',
                'Workflow smoke test on published auxiliary gamma-ray record; not reproduction of the paper primary Al-based astrochronology.',
                member='S4. Gamma-ray log_Albjära-1 core.xlsx'),
    ]
    print(json.dumps([{'file': r['output_file'], 'sha256': r['output_sha256'],
                       'rows': r['row_count']} for r in records], indent=2))


if __name__ == '__main__':
    main()
